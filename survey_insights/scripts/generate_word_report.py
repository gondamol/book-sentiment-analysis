#!/usr/bin/env python3
"""Generate an editable Word report and plot pack for the CLM survey analysis."""

from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
ROOT = Path(__file__).resolve().parents[1]
ANALYSIS_PATH = ROOT / "data" / "processed" / "analysis.json"
REPORT_DIR = ROOT / "reports"
FIGURE_DIR = REPORT_DIR / "figures"
DOCX_PATH = REPORT_DIR / "Tech_Oreon_Analytica_CLM_Feedback_Report.docx"
APP_DOCX_PATH = ROOT / "app" / "assets" / DOCX_PATH.name
APP_FIGURE_DIR = ROOT / "app" / "assets" / "figures"

NAVY = (16, 42, 67)
TEAL = (15, 139, 141)
RED = (199, 81, 70)
GOLD = (244, 185, 66)
SLATE = (107, 114, 128)
LINE = (217, 226, 236)
WHITE = (255, 255, 255)
BG = (248, 250, 252)


def load_analysis() -> dict:
    with ANALYSIS_PATH.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def ensure_dirs() -> None:
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    APP_FIGURE_DIR.mkdir(parents=True, exist_ok=True)


def make_pie_chart(title: str, items: list[dict], output_path: Path) -> None:
    width, height = 900, 520
    image = Image.new("RGB", (width, height), BG)
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((18, 18, width - 18, height - 18), radius=28, fill=WHITE, outline=LINE, width=2)
    draw.text((36, 32), title, fill=NAVY, font=find_font(28))

    total = sum(item["count"] for item in items) or 1
    box = (60, 110, 400, 450)
    palette = [TEAL, RED, NAVY, GOLD, SLATE]
    start = -90
    for index, item in enumerate(items):
        extent = (item["count"] / total) * 360
        draw.pieslice(box, start=start, end=start + extent, fill=palette[index % len(palette)], outline=WHITE, width=3)
        start += extent

    legend_x = 470
    legend_y = 130
    label_font = find_font(18)
    for index, item in enumerate(items):
        color = palette[index % len(palette)]
        draw.rounded_rectangle((legend_x, legend_y, legend_x + 28, legend_y + 28), radius=6, fill=color)
        pct = (item["count"] / total) * 100
        draw.text((legend_x + 44, legend_y - 1), f"{item['name']} ({pct:.1f}%)", fill=NAVY, font=label_font)
        legend_y += 52
    image.save(output_path)


def make_vertical_bar(title: str, labels: list[str], values: list[float], output_path: Path) -> None:
    width, height = 1100, 620
    image = Image.new("RGB", (width, height), BG)
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((18, 18, width - 18, height - 18), radius=28, fill=WHITE, outline=LINE, width=2)
    draw.text((36, 30), title, fill=NAVY, font=find_font(28))
    left, top, right, bottom = 90, 110, 1020, 500
    draw.line((left, bottom, right, bottom), fill=SLATE, width=3)
    draw.line((left, top, left, bottom), fill=SLATE, width=3)
    max_value = max(values) or 1
    step = max_value / 4
    small_font = find_font(16)
    label_font = find_font(18)
    for tick in range(5):
        y = bottom - ((bottom - top) * tick / 4)
        value = int(step * tick)
        draw.line((left - 6, y, right, y), fill=LINE, width=1)
        draw.text((24, y - 10), f"{value:,}", fill=SLATE, font=small_font)
    slot = (right - left) / len(values)
    bar_width = slot * 0.58
    for index, (label, value) in enumerate(zip(labels, values)):
        x0 = left + index * slot + (slot - bar_width) / 2
        x1 = x0 + bar_width
        bar_height = 0 if max_value == 0 else ((bottom - top) * value / max_value)
        y0 = bottom - bar_height
        draw.rounded_rectangle((x0, y0, x1, bottom), radius=10, fill=TEAL)
        draw.text((x0, y0 - 24), f"{int(value):,}", fill=NAVY, font=small_font)
        label_box = draw.textbbox((0, 0), label, font=label_font)
        label_x = x0 + (bar_width - (label_box[2] - label_box[0])) / 2
        draw.text((label_x, bottom + 12), label, fill=NAVY, font=label_font)
    image.save(output_path)


def make_horizontal_bar(title: str, labels: list[str], values: list[float], output_path: Path, color=RED) -> None:
    width, height = 1200, 700
    image = Image.new("RGB", (width, height), BG)
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((18, 18, width - 18, height - 18), radius=28, fill=WHITE, outline=LINE, width=2)
    draw.text((36, 30), title, fill=NAVY, font=find_font(28))
    left, top, right = 340, 110, 1080
    label_font = find_font(20)
    small_font = find_font(18)
    max_value = max(values) or 1
    row_gap = 78
    for index, (label, value) in enumerate(zip(labels, values)):
        y = top + index * row_gap
        draw.text((50, y + 10), label, fill=NAVY, font=label_font)
        draw.rounded_rectangle((left, y, right, y + 34), radius=10, fill=(240, 244, 248))
        bar_w = ((right - left) * value / max_value)
        draw.rounded_rectangle((left, y, left + bar_w, y + 34), radius=10, fill=color)
        draw.text((left + bar_w + 14, y + 5), f"{int(value):,}", fill=NAVY, font=small_font)
    image.save(output_path)


def make_line_chart(title: str, months: list[str], sat_values: list[float], access_values: list[float], output_path: Path) -> None:
    width, height = 1200, 680
    image = Image.new("RGB", (width, height), BG)
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((18, 18, width - 18, height - 18), radius=28, fill=WHITE, outline=LINE, width=2)
    draw.text((36, 30), title, fill=NAVY, font=find_font(28))
    left, top, right, bottom = 100, 120, 1080, 520
    draw.line((left, bottom, right, bottom), fill=SLATE, width=3)
    draw.line((left, top, left, bottom), fill=SLATE, width=3)
    small_font = find_font(16)
    label_font = find_font(18)
    for tick in range(6):
        y = bottom - ((bottom - top) * tick / 5)
        val = tick * 20
        draw.line((left, y, right, y), fill=LINE, width=1)
        draw.text((48, y - 10), f"{val}%", fill=SLATE, font=small_font)
    slot = (right - left) / max(len(months) - 1, 1)

    def points(series):
        pts = []
        for i, value in enumerate(series):
            x = left + i * slot
            y = bottom - ((bottom - top) * value / 100)
            pts.append((x, y))
        return pts

    sat_pts = points(sat_values)
    access_pts = points(access_values)
    if len(sat_pts) > 1:
        draw.line(sat_pts, fill=TEAL, width=4)
        draw.line(access_pts, fill=RED, width=4)
    for x, y in sat_pts:
        draw.ellipse((x - 5, y - 5, x + 5, y + 5), fill=TEAL)
    for x, y in access_pts:
        draw.ellipse((x - 5, y - 5, x + 5, y + 5), fill=RED)
    for i, month in enumerate(months):
        x = left + i * slot
        draw.text((x - 28, bottom + 16), month, fill=NAVY, font=label_font)
    draw.rounded_rectangle((760, 76, 1080, 126), radius=12, fill=(247, 249, 252), outline=LINE)
    draw.text((782, 88), "Positive satisfaction share", fill=TEAL, font=label_font)
    draw.text((782, 108), "Access challenge share", fill=RED, font=label_font)
    image.save(output_path)


def find_font(size: int):
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for candidate in candidates:
        path = Path(candidate)
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def make_word_cloud(title: str, terms: list[dict], output_path: Path, accent=(15, 139, 141)) -> None:
    width, height = 1400, 900
    image = Image.new("RGB", (width, height), (248, 250, 252))
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((18, 18, width - 18, height - 18), radius=28, fill=(255, 255, 255), outline=(222, 226, 230), width=2)
    title_font = find_font(42)
    draw.text((50, 45), title, fill=(16, 42, 67), font=title_font)

    if not terms:
        body_font = find_font(30)
        draw.text((50, 150), "No terms available", fill=(80, 91, 102), font=body_font)
        image.save(output_path)
        return

    max_count = max(item["count"] for item in terms) or 1
    palette = [
        accent,
        (199, 81, 70),
        (16, 42, 67),
        (244, 185, 66),
        (83, 105, 118),
    ]
    x = 50
    y = 140
    row_height = 60
    for index, item in enumerate(terms[:22]):
        size = int(28 + (item["count"] / max_count) * 44)
        font = find_font(size)
        word = item["term"]
        bbox = draw.textbbox((0, 0), word, font=font)
        word_width = bbox[2] - bbox[0]
        word_height = bbox[3] - bbox[1]
        if x + word_width > width - 70:
            x = 50
            y += row_height
            row_height = 60
        if y + word_height > height - 80:
            break
        draw.text((x, y), word, fill=palette[index % len(palette)], font=font)
        x += word_width + 24
        row_height = max(row_height, word_height + 20)
    image.save(output_path)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def make_summary_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Indicator"
    hdr[1].text = "Value"
    for cell in hdr:
        set_cell_shading(cell, "102A43")
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.bold = True
    for left, right in rows:
        row = table.add_row().cells
        row[0].text = left
        row[1].text = right


def add_caption(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(83, 105, 118)


def build_docx(analysis: dict, figures: dict) -> None:
    document = Document()
    core = document.core_properties
    core.author = "Tech Oreon Analytica"
    core.title = "HIV Care and Treatment Services: CLM Feedback Analysis"
    core.subject = "Editable report with charts and word clouds"

    styles = document.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)

    cover = document.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cover.add_run("TECH OREON ANALYTICA\n")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(16, 42, 67)
    run = cover.add_run("\nHIV Care and Treatment Services: CLM Feedback Analysis\n")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(16, 42, 67)
    run = cover.add_run("\nEditable stakeholder report\n")
    run.italic = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(83, 105, 118)
    run = cover.add_run("\nPrepared by Tech Oreon Analytica")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(15, 139, 141)

    document.add_paragraph("")
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(
        "This report presents findings from patient and community feedback on HIV care and treatment services, with a focus on service experience, satisfaction, access barriers, and priorities for improvement."
    )
    document.add_page_break()

    document.add_heading("Overview", level=1)
    document.add_paragraph(
        "The workbook contains 23,457 responses, dominated by CLM feedback and concentrated in Homa Bay, Kilifi, and Makueni. Overall experience is broadly positive, especially around staff attitude and reception, but a meaningful share of respondents still report access and service-delivery bottlenecks."
    )
    document.add_paragraph(
        "The analysis is presented in a CLM-style format so that the evidence can be used in stakeholder meetings, quality-improvement discussions, and routine review of what patients say about HIV services."
    )

    document.add_heading("Key Indicators", level=1)
    make_summary_table(
        document,
        [
            ("Total responses", f"{analysis['response_count']:,}"),
            ("Satisfied or very satisfied", f"{analysis['metrics']['positive_satisfaction_share']}%"),
            ("Reported access challenges", f"{analysis['metrics']['access_challenge_share']}%"),
            ("Confidentiality respected", f"{analysis['metrics']['confidentiality_yes_share']}%"),
            ("Rights awareness", f"{analysis['metrics']['rights_awareness_share']}%"),
            ("Average sentiment score", f"{analysis['metrics']['avg_sentiment']}"),
        ],
    )

    document.add_heading("Background and Context", level=1)
    document.add_paragraph(
        "This analysis sits within the broader HIV response ecosystem in Kenya, where community networks, civil society, government, and global health partners all play a role in improving access, treatment continuity, and quality of care. In this setting, community-led monitoring evidence is useful because it shows how services are actually experienced by the people who use them."
    )
    document.add_paragraph(
        "Organizations active in this wider ecosystem include NEPHAK, county and national health teams, PEPFAR/USAID, the Global Fund, UNAIDS, WHO, and other implementing and community partners supporting HIV prevention, treatment, care, and community systems."
    )

    document.add_heading("Plots and Findings", level=1)
    plot_sections = [
        (
            "Satisfaction and Sentiment",
            [figures["satisfaction_mix"], figures["sentiment_mix"]],
            "The feedback is largely positive, but the comments help explain where operational gaps continue to affect the experience of care.",
        ),
        (
            "Response Distribution",
            [figures["county_volume"], figures["survey_mix"]],
            "Most responses come from a few counties and are heavily concentrated in CLM feedback, which should be kept in mind when interpreting the profile of issues raised.",
        ),
        (
            "Service Priorities",
            [figures["theme_priority"], figures["service_trend"]],
            "The most frequently discussed themes point to concrete service-delivery issues rather than abstract satisfaction alone.",
        ),
        (
            "Client Language",
            [figures["positive_wordcloud"], figures["improvement_wordcloud"]],
            "The word clouds show the language respondents use most often when describing what worked well and what should improve.",
        ),
    ]

    for heading, image_paths, note in plot_sections:
        document.add_heading(heading, level=2)
        document.add_paragraph(note)
        for image_path in image_paths:
            document.add_picture(str(image_path), width=Inches(6.6))
            add_caption(document, image_path.stem.replace("_", " ").title())

    document.add_heading("Main Insights", level=1)
    for text in [
        "Staff friendliness, respectful reception, and supportive care are among the strongest positive themes in the feedback.",
        "The most common improvement requests relate to medicines and supplies, food and transport support, equipment and infrastructure, and waiting time.",
        "High satisfaction scores do not eliminate the need to act on practical barriers, because many respondents still describe specific service gaps that affect access and continuity of care.",
        "The open comments provide useful explanatory detail that can strengthen routine review meetings and help interpret service-quality indicators more clearly.",
    ]:
        document.add_paragraph(text, style="List Bullet")

    document.add_heading("Selected Client Voice", level=1)
    for quote in analysis["quotes"]["positive"][:2] + analysis["quotes"]["improvement"][:2]:
        p = document.add_paragraph()
        p.add_run(f"{quote['county']} | {quote['survey_name']}: ").bold = True
        p.add_run(f"\"{quote['snippet']}\"")

    document.add_heading("Recommendations", level=1)
    for text in [
        "Protect the positive service culture reflected in comments on friendly staff, respectful treatment, and supportive care.",
        "Prioritize action on medicines and supplies, food and transport support, equipment, and waiting time because these are the most consistent operational concerns in the feedback.",
        "Use both structured indicators and open comments in review meetings so that service-quality trends are interpreted alongside direct patient voice.",
        "Refresh these charts regularly as new feedback is collected so the report can support ongoing quality improvement rather than one-off presentation use.",
    ]:
        document.add_paragraph(text, style="List Number")

    for section in document.sections:
        header = section.header.paragraphs[0]
        header.text = "Tech Oreon Analytica"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(83, 105, 118)

    document.save(str(DOCX_PATH))
    APP_DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(DOCX_PATH, APP_DOCX_PATH)


def main() -> None:
    ensure_dirs()
    analysis = load_analysis()

    satisfaction_items = [item for item in analysis["satisfaction_counts"] if item["name"] != "Unknown"][:4]
    sentiment_items = analysis["sentiment_counts"][:3]
    county_items = analysis["county_counts"][:6]
    survey_items = analysis["survey_counts"][:5]
    theme_items = analysis["theme_overview"][:6]
    trend_items = [item for item in analysis["monthly_trend"] if item["responses"] >= 100 and "2022-10" <= item["month"] <= "2023-04"]

    figures = {
        "satisfaction_mix": FIGURE_DIR / "satisfaction_mix.png",
        "sentiment_mix": FIGURE_DIR / "sentiment_mix.png",
        "county_volume": FIGURE_DIR / "county_volume.png",
        "survey_mix": FIGURE_DIR / "survey_mix.png",
        "theme_priority": FIGURE_DIR / "theme_priority.png",
        "service_trend": FIGURE_DIR / "service_trend.png",
        "positive_wordcloud": FIGURE_DIR / "positive_wordcloud.png",
        "improvement_wordcloud": FIGURE_DIR / "improvement_wordcloud.png",
    }

    make_pie_chart("Satisfaction mix", satisfaction_items, figures["satisfaction_mix"])
    make_pie_chart("Sentiment mix", sentiment_items, figures["sentiment_mix"])
    make_vertical_bar(
        "Top counties by response volume",
        [item["name"] for item in county_items],
        [item["count"] for item in county_items],
        figures["county_volume"],
    )
    make_horizontal_bar(
        "Survey modules represented",
        [item["name"] for item in survey_items],
        [item["count"] for item in survey_items],
        figures["survey_mix"],
        color=TEAL,
    )
    make_horizontal_bar(
        "Most common themes in open comments",
        [item["label"] for item in theme_items],
        [item["count"] for item in theme_items],
        figures["theme_priority"],
        color=RED,
    )
    make_line_chart(
        "Positive satisfaction and access challenge trend",
        [item["month"] for item in trend_items],
        [item["positive_satisfaction_share"] for item in trend_items],
        [item["access_challenge_share"] for item in trend_items],
        figures["service_trend"],
    )
    make_word_cloud("Positive language", analysis["top_terms"]["positive"], figures["positive_wordcloud"], accent=(15, 139, 141))
    make_word_cloud("Improvement language", analysis["top_terms"]["improvements"], figures["improvement_wordcloud"], accent=(199, 81, 70))

    for image_path in figures.values():
        shutil.copyfile(image_path, APP_FIGURE_DIR / image_path.name)

    build_docx(analysis, figures)
    print(f"Wrote editable report: {DOCX_PATH}")
    print(f"Wrote figure pack to: {FIGURE_DIR}")
    print(f"Copied Word report to: {APP_DOCX_PATH}")


if __name__ == "__main__":
    main()
